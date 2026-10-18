from pathlib import Path
import openpyxl
p=Path(__file__).resolve().parents[2]/'qa'
s=openpyxl.load_workbook(p/'all.xlsx').worksheets[0]
rows={s.cell(r,3).value:r for r in range(7,s.max_row+1)}
r=rows['00123']; pending=rows['00002']
assert s.cell(r,3).data_type=='s'
assert s.cell(r,8).value==-2.625
assert s.cell(r,10).data_type=='s'
assert s.cell(pending,7).value in ('',None)
assert '12345678901234567890' in rows
assert s.freeze_panes=='A7' and s.sheet_view.rightToLeft
assert s.auto_filter.ref=='A6:L9'
assert s['A1'].fill.fgColor.rgb=='FF183B43'
assert 'محلية' in s['A3'].value and 'مرحلي' in s['A3'].value
assert s.cell(r,8).number_format=='#,##0.############;[Red](#,##0.############);0'
assert openpyxl.load_workbook(p/'diff.xlsx').worksheets[0].max_row==7
assert openpyxl.load_workbook(p/'empty.xlsx').worksheets[0]['A6'].value=='المجموعة'
print('PASS: independent Excel readback, exact values and formatting')
